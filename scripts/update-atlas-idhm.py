#!/usr/bin/env python3
"""Gera snapshot municipal dos componentes do IDHM a partir da base oficial Atlas Brasil.

Uso:
  python scripts/update-atlas-idhm.py

Dependência de manutenção:
  pip install openpyxl

A fonte oficial atualmente publicada pelo Atlas Brasil apresenta um certificado TLS
que pode não ser aceito por alguns runners. Por padrão a verificação TLS permanece
ligada. Para uma atualização auditada, após confirmar manualmente o domínio e o
caminho, use ATLAS_INSECURE_TLS=1.
"""

from __future__ import annotations

import hashlib
import json
import os
import ssl
import tempfile
import urllib.request
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

SOURCE_URL = (
    "https://atlasbrasil.org.br/cockpit/storage/uploads/dados/"
    "censo_total_1991_2010.xlsx"
)
SHEET = "MUN 91-00-10"
EXPECTED_YEARS = [1991, 2000, 2010]
EXPECTED_MUNICIPALITIES = 5565
REQUIRED = ["ANO", "Codmun7", "IDHM", "IDHM_E", "IDHM_L", "IDHM_R"]
OUT = Path("src/data/atlas-idhm-municipios-1991-2010.json")


def download(url: str, path: Path) -> None:
    context = None
    if os.getenv("ATLAS_INSECURE_TLS") == "1":
        context = ssl._create_unverified_context()
    with urllib.request.urlopen(url, context=context) as src, path.open("wb") as dst:
        while chunk := src.read(1024 * 1024):
            dst.write(chunk)


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def main() -> None:
    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "censo_total_1991_2010.xlsx"
        download(SOURCE_URL, xlsx)

        wb = load_workbook(xlsx, read_only=True, data_only=True)
        if SHEET not in wb.sheetnames:
            raise RuntimeError(f"Aba esperada ausente: {SHEET}")

        ws = wb[SHEET]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        idx = {str(value): i for i, value in enumerate(header) if value is not None}
        missing = [name for name in REQUIRED if name not in idx]
        if missing:
            raise RuntimeError(f"Colunas esperadas ausentes: {missing}")

        municipalities: dict[str, dict[str, list[float]]] = {}
        row_count = 0

        for row in rows:
            year = int(row[idx["ANO"]])
            if year not in EXPECTED_YEARS:
                continue
            code = str(int(row[idx["Codmun7"]]))
            values = [
                float(row[idx["IDHM"]]),
                float(row[idx["IDHM_E"]]),
                float(row[idx["IDHM_L"]]),
                float(row[idx["IDHM_R"]]),
            ]
            municipalities.setdefault(code, {})[str(year)] = values
            row_count += 1

        if len(municipalities) != EXPECTED_MUNICIPALITIES:
            raise RuntimeError(
                f"Municípios: esperado {EXPECTED_MUNICIPALITIES}, obtido {len(municipalities)}"
            )
        expected_rows = EXPECTED_MUNICIPALITIES * len(EXPECTED_YEARS)
        if row_count != expected_rows:
            raise RuntimeError(f"Registros: esperado {expected_rows}, obtido {row_count}")

        incomplete = [
            code
            for code, years in municipalities.items()
            if sorted(map(int, years.keys())) != EXPECTED_YEARS
        ]
        if incomplete:
            raise RuntimeError(f"Municípios sem os três anos: {len(incomplete)}")

        payload = {
            "schema_version": 1,
            "source": "Atlas do Desenvolvimento Humano no Brasil",
            "source_url": SOURCE_URL,
            "source_sheet": SHEET,
            "source_sha256": sha256(xlsx),
            "retrieved_at": datetime.now(timezone.utc).isoformat(),
            "reference_years": EXPECTED_YEARS,
            "value_order": ["IDHM", "IDHM_E", "IDHM_L", "IDHM_R"],
            "municipality_count": len(municipalities),
            "record_count": row_count,
            "municipalities": dict(sorted(municipalities.items())),
        }

        OUT.parent.mkdir(parents=True, exist_ok=True)
        OUT.write_text(
            json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + "\n",
            encoding="utf-8",
        )
        print(
            json.dumps(
                {
                    "output": str(OUT),
                    "municipalities": len(municipalities),
                    "records": row_count,
                    "source_sha256": payload["source_sha256"],
                },
                ensure_ascii=False,
            )
        )


if __name__ == "__main__":
    main()
